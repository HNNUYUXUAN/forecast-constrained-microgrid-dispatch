"""Tables, figures, and portable template exports for the dispatch study."""
from __future__ import annotations
from copy import copy
from datetime import datetime, timedelta
import gzip
import json
import os
from pathlib import Path
import sys
import lzma
import struct
import hashlib
from functools import lru_cache

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from c_grid.src.online_dispatch import COLUMNS,DT,summarize

NAMES={'q2':'日前计划·固定价格','q3':'预报更新·固定价格','q42':'日前计划·价格预测',
       'q43':'预报更新·价格预测','q42_announced':'日前计划·价格公布','q43_announced':'预报更新·价格公布'}
FILES={'q2':'result2.xlsx','q3':'result3.xlsx','q42':'result4-2.xlsx','q43':'result4-3.xlsx',
       'q42_announced':'result4-2.xlsx','q43_announced':'result4-3.xlsx'}


def load_records(run,name):
    if not (Path(run)/name/'intervals.npz').exists():
        index,blocks=read_compact_records(str(Path(run)/'process_records.xz'))
        shape=index[name]['shape']
        return np.stack([np.frombuffer(blocks[k],dtype=np.float64).reshape(shape[:2]) for k in index[name]['columns']],axis=2)
    with np.load(Path(run)/name/'intervals.npz') as saved:
        return saved['records']


@lru_cache(maxsize=1)
def read_compact_records(path):
    with lzma.open(path,'rb') as f:
        length=struct.unpack('<Q',f.read(8))[0]
        assert length<1000000
        header=json.loads(f.read(length))
        blocks={}
        for key,length in header['blocks']:
            shuffled=f.read(length)
            raw=np.frombuffer(shuffled,dtype=np.uint8).reshape(8,-1).T.copy().tobytes()
            assert hashlib.sha256(raw).hexdigest()==key
            blocks[key]=raw
        assert not f.read(1)
    return header['index'],blocks


def date_for(day):
    return datetime(2025,1,1)+timedelta(days=int(day))


def clock(minutes):
    if minutes==1440:
        return '24:00'
    return f'{minutes//60:02d}:{minutes%60:02d}'


def emergency_events(values):
    events=[]
    i=0
    while i<144:
        if values[i]*DT<=1e-8:
            i+=1
            continue
        start=i
        while i<144 and values[i]*DT>1e-8:
            i+=1
        events.append((f'{clock(start*10)}-{clock(i*10)}',float(values[start:i].sum()*DT)))
    return events


def storage_rows(charge,discharge,soc_start,soc_end,date=None):
    rows=[]
    for block in range(6):
        period=f'{block*4}:00-{(block+1)*4}:00'
        when='0:00' if block==0 else '24:00' if block==1 else None
        level=float(soc_start) if block==0 else float(soc_end) if block==1 else None
        row=[period,float(charge[block*24:(block+1)*24].sum()*DT),
             float(discharge[block*24:(block+1)*24].sum()*DT),when,level]
        if date is not None:
            row.insert(0,date)
        rows.append(row)
    return rows


def workbook_payload(run):
    run=Path(run)
    payload={}
    q1=pd.read_csv(run/'problem1.csv')
    template=load_workbook(ROOT/'c_grid/attachments/附件5/result1.xlsx',read_only=True)
    labels=[r[0] for r in list(template['计划购电量'].values)[1:]]
    template.close()
    payload['results/result1.xlsx']={'计划购电量':[[label,float(x)] for label,x in zip(labels,q1.grid_purchase_kwh)],
       '充放电量':storage_rows(q1.charge_kwh.to_numpy()/DT,q1.discharge_kwh.to_numpy()/DT,
                          q1.storage_start_kwh.iloc[0],q1.storage_end_kwh.iloc[-1])}
    for name,filename in FILES.items():
        records=load_records(run,name)
        planned,adjusted,storage,emergency=[],[],[],[]
        for index,row in enumerate(records):
            date=date_for(index+31).date().isoformat()
            a={key:row[:,i] for i,key in enumerate(COLUMNS)}
            summary=summarize(row)
            # The plan sheet records the zero-hour contractual energy and its tariff cost.
            planned.append([date,*a['plan_kw']*DT,float(a['plan_kw'].sum()*DT),summary['planned_cost_yuan']])
            # The adjusted sheet records effective energy and the full realized daily cost.
            adjusted.append([date,*a['effective_kw']*DT,float(a['effective_kw'].sum()*DT),summary['total_cost_yuan']])
            storage.extend(storage_rows(a['charge_kw'],a['discharge_kw'],a['soc_start_kwh'][0],a['soc_end_kwh'][-1],date))
            for span,amount in emergency_events(a['emergency_kw']):
                emergency.append([date,span,amount])
        sheets={'计划购电量':planned}
        if name.startswith(('q3','q43')):
            sheets['调整购电量']=adjusted
        sheets.update({'充放电量':storage,'紧急购电量':emergency})
        folder='comparison_announced' if name.endswith('announced') else 'results'
        payload[f'{folder}/{filename}']=sheets
    return payload


def export_workbooks_portable(run,output=None):
    """Python template writer for environments without the desktop spreadsheet runtime."""
    run=Path(run)
    output=Path(output or run)
    payload=workbook_payload(run)
    for relative,sheets in payload.items():
        target=output/relative
        target.parent.mkdir(parents=True,exist_ok=True)
        workbook=load_workbook(ROOT/'c_grid/attachments/附件5'/target.name)
        workbook.properties.creator=''
        workbook.properties.lastModifiedBy=''
        for name,rows in sheets.items():
            sheet=workbook[name]
            styles=[copy(sheet.cell(2,c)._style) for c in range(1,sheet.max_column+1)]
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.value=None
            if len(rows)+1<sheet.max_row:
                sheet.delete_rows(len(rows)+2,sheet.max_row-len(rows)-1)
            for r,values in enumerate(rows,2):
                for c,value in enumerate(values,1):
                    if c==1 and isinstance(value,str) and len(value)==10 and value[4]=='-' and value[7]=='-':
                        value=datetime.fromisoformat(value)
                    cell=sheet.cell(r,c,value.item() if isinstance(value,np.generic) else value)
                    cell._style=copy(styles[c-1])
        workbook.save(target)
    with gzip.open(output/'excel_payload.json.gz','wt',encoding='utf-8') as handle:
        json.dump(payload,handle,ensure_ascii=False,default=float)
    return payload


def verify_workbooks(run,output=None):
    output=Path(output or run)
    payload=workbook_payload(run)
    checked=0
    for relative,sheets in payload.items():
        workbook=load_workbook(output/relative,read_only=True,data_only=True)
        template=load_workbook(ROOT/'c_grid/attachments/附件5'/Path(relative).name,read_only=True)
        assert workbook.sheetnames==template.sheetnames
        for name,rows in sheets.items():
            sheet=workbook[name]
            assert list(next(sheet.values))==list(next(template[name].values)),(relative,name,'header')
            actual=list(sheet.values)[1:]
            assert len(actual)==len(rows),(relative,name,len(actual),len(rows))
            for r,(got,want) in enumerate(zip(actual,rows)):
                assert len(got)==len(want)
                for left,right in zip(got,want):
                    if isinstance(right,(int,float,np.number)):
                        assert left is not None and abs(left-right)<1e-6,(relative,name,r,left,right)
                    elif isinstance(right,str) and len(right)==10 and right[4]=='-' and right[7]=='-':
                        assert left==datetime.fromisoformat(right)
                    else:
                        assert left==right,(relative,name,r,left,right)
                    checked+=1
        workbook.close()
        template.close()
    return {'workbooks':len(payload),'checked_cells':checked,'status':'PASS'}


def specified_tables(run,name):
    records=load_records(run,name)
    purchase,storage,emergency=[],[],[]
    for date in ('2025-03-20','2025-06-21','2025-09-23','2025-12-21'):
        day=(datetime.fromisoformat(date)-datetime(2025,2,1)).days
        row=records[day]
        a={key:row[:,i] for i,key in enumerate(COLUMNS)}
        summary=summarize(row)
        p={'日期':date}
        for hour in (10,12,14,16,18,20):
            p[f'{hour}:00-{hour}:10 / kWh']=a['effective_kw'][hour*6]*DT
        p.update({'全天购电量 / kWh':a['effective_kw'].sum()*DT,'全天总费用 / 元':summary['total_cost_yuan']})
        purchase.append(p)
        for values in storage_rows(a['charge_kw'],a['discharge_kw'],a['soc_start_kwh'][0],a['soc_end_kwh'][-1],date):
            storage.append(values)
        for span,amount in emergency_events(a['emergency_kw']):
            emergency.append([date,span,amount])
    return (pd.DataFrame(purchase),pd.DataFrame(storage,columns=['日期','时段','充电量 / kWh','放电量 / kWh','时刻','储电量 / kWh']),
            pd.DataFrame(emergency,columns=['日期','时段','紧急购电量 / kWh']))


def configure_style():
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import font_manager as fm
    locations=[Path(os.environ.get('MICROGRID_FONT_DIR','fonts')),Path.home()/'.local/share/fonts/microgrid',Path('C:/Windows/Fonts')]
    found={}
    for directory in locations:
        for name in ('simsun.ttc','times.ttf','timesbd.ttf','timesi.ttf'):
            path=directory/name
            if path.is_file():
                fm.fontManager.addfont(str(path))
                found[name]=path
    for family in ('SimSun','Times New Roman'):
        fm.findfont(fm.FontProperties(family=family),fallback_to_default=False)
    matplotlib.rcParams.update({'font.family':['Times New Roman','SimSun'],'font.size':11,
        'axes.unicode_minus':False,'axes.spines.top':False,'axes.spines.right':False,
        'axes.prop_cycle':matplotlib.cycler(color=['black','#444444','#777777','#aaaaaa']),
        'figure.dpi':120,'savefig.dpi':200,'mathtext.fontset':'stix','svg.fonttype':'path'})


def create_figures(run):
    configure_style()
    import matplotlib.pyplot as plt
    run=Path(run)
    directory=run/'figures'
    directory.mkdir(exist_ok=True)
    summary=pd.DataFrame(json.loads((run/'summary.json').read_text())).set_index('name')
    paths=[]
    def save(fig,name):
        fig.tight_layout()
        fig.savefig(directory/(name+'.png'),bbox_inches='tight')
        fig.savefig(directory/(name+'.svg'),bbox_inches='tight')
        plt.close(fig)
        paths.append(name)
    fig,ax=plt.subplots(figsize=(10,4.5))
    order=list(NAMES)
    base=np.zeros(len(order))
    for field,label,hatch in [('planned_cost_yuan','零点计划费用',''),('incremental_cost_yuan','调整增购费用','///'),('emergency_cost_yuan','紧急购电费用','xxx')]:
        values=summary.loc[order,field].to_numpy()/1e4
        ax.bar(np.arange(len(order)),values,bottom=base,label=label,color='white',edgecolor='black',hatch=hatch)
        base+=values
    credit=summary.loc[order,'cancellation_credit_yuan'].to_numpy()/1e4
    ax.scatter(np.arange(len(order)),summary.loc[order,'total_cost_yuan']/1e4,marker='D',c='black',label='扣除取消返还后的总费用',zorder=4)
    ax.set_xticks(np.arange(len(order)),[NAMES[n].replace('·','\n') for n in order])
    ax.set_ylabel('费用 / 万元')
    ax.legend(fontsize=9,ncol=2)
    save(fig,'cost_components')
    records=load_records(run,'q3')
    fig,axes=plt.subplots(4,2,figsize=(12,12),sharex=True)
    for pair,date in zip(axes,('2025-03-20','2025-06-21','2025-09-23','2025-12-21')):
        i=(datetime.fromisoformat(date)-datetime(2025,2,1)).days
        row=records[i]; x=np.arange(144)/6
        pair[0].plot(x,row[:,9]-row[:,10],label='实际净负荷',color='black')
        pair[0].plot(x,row[:,0],label='零点计划',color='gray',ls='--')
        pair[0].plot(x,row[:,1],label='有效购电',color='black',ls=':')
        pair[0].set(title=date,ylabel='功率 / kW')
        pair[1].plot(np.arange(145)/6,np.r_[row[0,6],row[:,7]],color='black',label='储电量')
        pair[1].axhline(1200,color='gray',ls='--');pair[1].axhline(10800,color='gray',ls='--')
        pair[1].set(title=date,ylabel='储电量 / kWh',ylim=(0,12000))
    axes[0,0].legend(fontsize=9);axes[0,1].legend(fontsize=9)
    for ax in axes[-1]: ax.set_xlabel('时刻 / h')
    save(fig,'specified_dispatch')
    fig,axes=plt.subplots(2,1,figsize=(10,7),sharex=True)
    for name,line in [('q2','--'),('q3','-'),('q42',':'),('q43','-.')]:
        daily=pd.DataFrame(json.loads((run/name/'daily.json').read_text()))
        daily['month']=[date_for(d).month for d in daily.day]
        grouped=daily.groupby('month')[['total_cost_yuan','emergency_kwh']].sum()
        axes[0].plot(grouped.index,grouped.total_cost_yuan/1e4,ls=line,color='black',label=NAMES[name])
        axes[1].plot(grouped.index,grouped.emergency_kwh,ls=line,color='black',label=NAMES[name])
    axes[0].set_ylabel('月费用 / 万元');axes[1].set_ylabel('紧急购电量 / kWh');axes[1].set_xlabel('月份')
    axes[0].legend(ncol=2,fontsize=9);axes[1].set_xticks(range(2,13))
    save(fig,'monthly_results')
    fig,ax=plt.subplots(figsize=(10,4.5))
    labels=[]
    for mask in range(8):
        bits=f'{mask:03b}'
        labels.append('0'+''.join('+'+str(h) for h,b in zip((6,12,18),bits) if b=='1'))
    for prefix,line in [('q3','-'),('q43','--'),('q43_announced',':')]:
        values=[summary.loc[prefix if mask==7 else prefix+'_updates_'+f'{mask:03b}','total_cost_yuan']/1e4 for mask in range(8)]
        ax.plot(range(8),values,marker='o',ls=line,color='black',label=NAMES[prefix])
    ax.set_xticks(range(8),[s.replace('+','\n+') if len(s)>6 else s for s in labels])
    ax.set(xlabel='采用的预报发布时间 / h',ylabel='费用 / 万元');ax.legend(fontsize=9)
    save(fig,'forecast_updates')
    fig,axes=plt.subplots(1,2,figsize=(12,4.5))
    for ax,prefix in zip(axes,('q3','q43')):
        names=[prefix,prefix+'_risk_0.5',prefix+'_risk_0.9',prefix+'_terminal_0.5',prefix+'_terminal_1.5',prefix+'_daily_closure',prefix+'_constant']
        delta=(summary.loc[names,'total_cost_yuan']/summary.loc[prefix,'total_cost_yuan']-1)*100
        ax.barh(range(len(names)),delta,color='white',edgecolor='black',hatch='///')
        ax.set_yticks(range(len(names)),['主配置','分位数0.5','分位数0.9','终端价值0.5倍','终端价值1.5倍','日闭环','分段常值预报'])
        ax.set(title=NAMES[prefix],xlabel='相对主配置费用变化 / %');ax.axvline(0,color='black',lw=.8)
    save(fig,'sensitivity')
    fig,axes=plt.subplots(1,2,figsize=(10,4))
    for ax,name in zip(axes,('q3','q43')):
        a=load_records(run,name)
        errors=a[:,:,13]-(a[:,:,9]-a[:,:,10])
        ax.boxplot([errors[:,i:i+36].ravel() for i in (0,36,72,108)],tick_labels=['0-6','6-12','12-18','18-24'],showfliers=False)
        ax.axhline(0,color='gray',ls='--');ax.set(title=NAMES[name],xlabel='日内时段 / h',ylabel='风险调整净负荷预测误差 / kW')
    save(fig,'forecast_errors')
    return paths


def main():
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('--run',type=Path,required=True)
    parser.add_argument('--export',action='store_true')
    parser.add_argument('--figures',action='store_true')
    args=parser.parse_args()
    if args.export:
        export_workbooks_portable(args.run)
        print(json.dumps(verify_workbooks(args.run),ensure_ascii=False))
    if args.figures:
        print(create_figures(args.run))


if __name__=='__main__':main()
